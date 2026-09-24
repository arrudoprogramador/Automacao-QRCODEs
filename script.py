import glob
import os
import re
import unicodedata
from difflib import SequenceMatcher

import openpyxl
from PyPDF2 import PdfReader, PdfWriter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PASTA_CONFIDENCIAIS = os.path.join(BASE_DIR, "confidenciais")
PASTA_PDFS = os.path.join(PASTA_CONFIDENCIAIS, "pdfs")
PASTA_PLANILHAS = os.path.join(PASTA_CONFIDENCIAIS, "planilha")
PASTA_SAIDA = os.path.join(PASTA_CONFIDENCIAIS, "separados")
CAMINHO_RELATORIO = os.path.join(PASTA_CONFIDENCIAIS, "relatorio_nao_encontrados.txt")

RATIO_CONFIANCA = 0.85
MARGEM_VENCEDOR = 0.08


def normalizar(texto):
    """Remove acentos, pontuação, espaços duplicados e tudo para minúsculas."""
    texto = unicodedata.normalize("NFD", str(texto))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    texto = texto.lower()
    texto = re.sub(r"[^a-z0-9 ]", " ", texto)
    return " ".join(texto.split())


def razao_similaridade(a, b):
    return SequenceMatcher(None, normalizar(a), normalizar(b)).ratio()


def extrair_nome_paciente(texto_pagina):
    m = re.search(r"Paciente:\s*(.+)", texto_pagina or "")
    return m.group(1).strip() if m else ""


def extrair_data_exame(texto_pagina):
    m = re.search(r"Data do exame:\s*(\d{2})/(\d{2})/(\d{4})", texto_pagina or "")
    return m.groups() if m else None  # (dd, mm, aaaa)


def ler_chave_dia_planilha(caminho_planilha, nome_aba=None):
    """Devolve a chave 'dia-DD-MM' lendo a linha 'Data: DD/MM/AAAA' da planilha."""
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    ws = wb[nome_aba] if nome_aba else wb.active
    for linha in ws.iter_rows(values_only=True):
        texto = str(linha[0]) if linha and linha[0] is not None else ""
        m = re.search(r"Data:\s*(\d{2})/(\d{2})/(\d{4})", texto)
        if m:
            dia, mes, _ = m.groups()
            return f"dia-{dia}-{mes}", (dia, mes)
    raise ValueError(f"Data DD/MM/AAAA não encontrada em {caminho_planilha} (aba {nome_aba or 'ativa'})")


def ler_pacientes_planilha(caminho_planilha, nome_aba=None):
    """Lê a aba da planilha encontrando automaticamente a linha de cabeçalho (Nº/Nome)."""
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    ws = wb[nome_aba] if nome_aba else wb.active
    linhas = list(ws.iter_rows(values_only=True))

    cabecalho = None
    for idx, linha in enumerate(linhas):
        rotulos = [str(c).strip().lower() if c is not None else "" for c in linha]
        if "nome" in rotulos and any(r in ("nº", "num", "n.") for r in rotulos):
            cabecalho = idx
            break
    if cabecalho is None:
        raise ValueError(f"Cabeçalho 'Nº'/'Nome' não encontrado em {caminho_planilha}")

    col_nome = next(i for i, c in enumerate(linhas[cabecalho]) if str(c).lower().strip() == "nome")
    col_id = next(
        i for i, c in enumerate(linhas[cabecalho])
        if str(c).lower().strip() in ("nº", "num", "n.")
    )

    pacientes = []
    for linha in linhas[cabecalho + 1:]:
        nome = linha[col_nome]
        ide = linha[col_id]
        if nome is None or ide is None:
            continue
        try:
            ide = int(float(ide))
        except (ValueError, TypeError):
            continue
        nome = str(nome).strip()
        if not nome or nome.lower() == "total":
            continue
        pacientes.append((ide, nome))
    return pacientes


def carregar_paginas(caminhos_pdf):
    """Agrupa páginas pela data do exame: {chave_dia: [(caminho, num, nome_norm, nome_orig)]}."""
    paginas_por_dia = {}
    for caminho in caminhos_pdf:
        leitor = PdfReader(caminho)
        for num in range(len(leitor.pages)):
            texto = leitor.pages[num].extract_text()
            data = extrair_data_exame(texto)
            if not data:
                continue
            chave = f"dia-{data[0]}-{data[1]}"
            nome_original = extrair_nome_paciente(texto)
            paginas_por_dia.setdefault(chave, []).append(
                (caminho, num, normalizar(nome_original), nome_original)
            )
    return paginas_por_dia


def casar_nomes(nome_planilha, paginas):
    """
    Retorna a melhor página candidata se a correspondência for de alta confiança.
    Confiança: normalizados iguais OU (ratio >= RATIO_CONFIANCA E folga do 2º >= MARGEM).
    """
    alvo = normalizar(nome_planilha)
    melhores = sorted(
        ((razao_similaridade(nome_planilha, pagina[3]), pagina)
         for pagina in paginas if pagina[2]),
        reverse=True,
    )
    if not melhores:
        return None

    melhor_ratio, melhor_pagina = melhores[0]
    segundo_ratio = melhores[1][0] if len(melhores) > 1 else 0.0

    if alvo == melhor_pagina[2]:
        return melhor_pagina
    if melhor_ratio >= RATIO_CONFIANCA and melhor_ratio - segundo_ratio >= MARGEM_VENCEDOR:
        return melhor_pagina
    return None


def gerar_pdf_por_id(pagina, paciente_id, destino):
    caminho, numero_pagina, _, _ = pagina
    leitor = PdfReader(caminho)
    escritor = PdfWriter()
    escritor.add_page(leitor.pages[numero_pagina])
    with open(destino, "wb") as f:
        escritor.write(f)
    print(f"OK   : {paciente_id}.pdf criado a partir de {os.path.basename(caminho)} (pág. {numero_pagina + 1})")


def abas_com_dados(caminho_planilha):
    """Abas que contêm a linha 'Data: DD/MM/AAAA' seguida de cabeçalho Nº/Nome."""
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    abas = []
    for nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        linhas = list(ws.iter_rows(values_only=True))
        tem_data = any(
            re.search(r"Data:\s*\d{2}/\d{2}/\d{4}", str(linha[0] or ""))
            for linha in linhas
        )
        tem_cabecalho = any(
            "nome" in [str(c or "").strip().lower() for c in linha]
            and any(r in ("nº", "num", "n.") for r in [str(c or "").strip().lower() for c in linha])
            for linha in linhas
        )
        if tem_data and tem_cabecalho:
            abas.append(nome_aba)
    return abas


def processar(caminho_planilha=None, pasta_pdfs=None, pasta_saida=None):
    pasta_pdfs = pasta_pdfs or PASTA_PDFS
    pasta_saida = pasta_saida or PASTA_SAIDA

    caminhos_planilha = ([caminho_planilha] if caminho_planilha
                         else sorted(glob.glob(os.path.join(PASTA_PLANILHAS, "*.xlsx"))))
    if not caminhos_planilha:
        print("Nenhuma planilha encontrada em", PASTA_PLANILHAS)
        return

    caminhos_pdf = sorted(glob.glob(os.path.join(pasta_pdfs, "**", "*.pdf"), recursive=True))
    if not caminhos_pdf:
        print("Nenhum PDF encontrado em", pasta_pdfs)
        return

    paginas_por_dia = carregar_paginas(caminhos_pdf)

    gerados = 0
    nao_encontrados_por_dia = {}

    for caminho_planilha in caminhos_planilha:
        abas = abas_com_dados(caminho_planilha)
        if not abas:
            print(f"Nenhuma aba com 'Data:' + cabeçalho Nº/Nome em {caminho_planilha}")
            continue

        for nome_aba in abas:
            chave_dia, _ = ler_chave_dia_planilha(caminho_planilha, nome_aba)
            pacientes = ler_pacientes_planilha(caminho_planilha, nome_aba)
            paginas = paginas_por_dia.get(chave_dia, [])
            destino_dia = os.path.join(pasta_saida, chave_dia)
            os.makedirs(destino_dia, exist_ok=True)

            if not paginas:
                print(f"\n== {os.path.basename(caminho_planilha)} [{nome_aba}] ({chave_dia}): "
                      f"nenhuma página com essa data nos PDFs.")
                nao_encontrados_por_dia[chave_dia] = pacientes
                continue

            print(f"\n== {os.path.basename(caminho_planilha)} [{nome_aba}] ({chave_dia}): "
                  f"{len(pacientes)} pacientes, {len(paginas)} páginas disponíveis")

            utilizadas = set()
            nao_encontrados = []

            for paciente_id, nome in pacientes:
                destino = os.path.join(destino_dia, f"{paciente_id}.pdf")
                if os.path.exists(destino):
                    print(f"--  : {paciente_id}.pdf já existe, ignorado.")
                    continue

                candidata = casar_nomes(nome, paginas)
                chave = (candidata[0], candidata[1]) if candidata else None
                if candidata is None or chave in utilizadas:
                    nao_encontrados.append((paciente_id, nome))
                    print(f"AVISO: ID {paciente_id} ({nome}) não encontrado com alta confiança.")
                    continue

                utilizadas.add(chave)
                gerar_pdf_por_id(candidata, paciente_id, destino)
                gerados += 1

            nao_encontrados_por_dia[chave_dia] = nao_encontrados

    total_nao = sum(len(v) for v in nao_encontrados_por_dia.values())
    if total_nao:
        linhas = [f"Pacientes sem correspondência de alta confiança:", ""]
        for chave_dia, lista in nao_encontrados_por_dia.items():
            if not lista:
                continue
            linhas.append(f"== {chave_dia}")
            linhas += [f"  ID {paciente_id}: {nome}" for paciente_id, nome in lista]
            linhas.append("")
        with open(CAMINHO_RELATORIO, "w", encoding="utf-8") as f:
            f.write("\n".join(linhas) + "\n")
        print(f"\n{total_nao} paciente(s) sem correspondência de alta confiança "
              f"(detalhes em {os.path.basename(CAMINHO_RELATORIO)}):")
        for chave_dia, lista in nao_encontrados_por_dia.items():
            for paciente_id, nome in lista:
                print(f"  [{chave_dia}] ID {paciente_id}: {nome}")
    else:
        if os.path.exists(CAMINHO_RELATORIO):
            os.remove(CAMINHO_RELATORIO)

    print(f"\nGerados: {gerados} arquivo(s) em {pasta_saida}")


if __name__ == "__main__":
    processar()